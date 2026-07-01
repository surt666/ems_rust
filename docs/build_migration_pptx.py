#!/usr/bin/env python3
"""Build the counter-measurement migration presentation (.pptx) with native shapes.
Run: uv run --with python-pptx python3 build_migration_pptx.py
Content mirrors counter-measurement-migration-walkthrough.md; diagrams are editable shapes.
"""
# pyright: reportMissingImports=false
# (python-pptx is provided by `uv run --with python-pptx`, not the global env)
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- palette (muted, understated) ---
NAVY = RGBColor(0x1F, 0x3A, 0x4D)
BLUE = RGBColor(0x2E, 0x6E, 0x8E)
TEAL = RGBColor(0x3A, 0x8F, 0xA0)
GREEN = RGBColor(0x3B, 0x7A, 0x57)
AMBER = RGBColor(0xB0, 0x7D, 0x2A)
GRAY = RGBColor(0x8A, 0x8A, 0x8A)
LGRAY = RGBColor(0xDD, 0xE3, 0xE7)
LIGHT = RGBColor(0xEE, 0xF2, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x30, 0x3A)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def slide():
    return prs.slides.add_slide(BLANK)


def fill_text(tf, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [(paras, 14, INK, False)]
    for i, spec in enumerate(paras):
        text, size, color, bold = spec
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size); f.bold = bold; f.name = FONT; f.color.rgb = color
    return tf


def box(sl, x, y, w, h, paras, fill=LIGHT, line=None, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.0):
    sp = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    tf = sp.text_frame
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.04)
    fill_text(tf, paras, align, anchor)
    return sp


def arrow(sl, x, y, w, h, color=GRAY, shape=MSO_SHAPE.RIGHT_ARROW):
    sp = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    return sp


def textbox(sl, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(tb.text_frame, paras, align, anchor)
    return tb


def title_bar(sl, title):
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(0.95))
    bar.shadow.inherit = False
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.5)
    fill_text(tf, [(title, 26, WHITE, True)], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    accent = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.95), Inches(SW), Inches(0.06))
    accent.shadow.inherit = False
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()


def footer(sl, text):
    tb = sl.shapes.add_textbox(Inches(0.5), Inches(SH - 0.62), Inches(SW - 1.0), Inches(0.45))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "In short:  "
    r.font.size = Pt(13); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = GREEN
    r2 = p.add_run(); r2.text = text
    r2.font.size = Pt(13); r2.font.bold = False; r2.font.name = FONT; r2.font.color.rgb = INK


def bullets(sl, x, y, w, h, items, size=16, color=INK):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        indent = 0
        if isinstance(it, tuple):
            it, indent = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        p.level = indent
        r = p.add_run()
        r.text = ("   " * indent) + "•  " + it
        r.font.size = Pt(size - indent); r.font.name = FONT; r.font.color.rgb = color
    return tb


def down(sl, cx, y, h=0.32, w=0.34, color=LGRAY):
    arrow(sl, cx - w / 2, y, w, h, color=color, shape=MSO_SHAPE.DOWN_ARROW)


def right(sl, x, cy, w=0.3, h=0.34, color=LGRAY):
    arrow(sl, x, cy - h / 2, w, h, color=color, shape=MSO_SHAPE.RIGHT_ARROW)


# ============================================================ TITLE
s = slide()
bg = box(s, 0, 0, SW, SH, "", fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, 0, 2.55, SW, 0.06, "", fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
textbox(s, 0.9, 1.35, 11.5, 1.2,
        [("Counter-measurement migration", 40, WHITE, True)])
textbox(s, 0.9, 2.75, 11.5, 1.0,
        [("From the current services to the new data pipeline", 22, LGRAY, False)])
textbox(s, 0.9, 6.5, 11.5, 0.5,
        [("A walkthrough — detail in counter-measurement-migration.md", 14, LGRAY, False)])

# ============================================================ 1. TODAY
s = slide()
title_bar(s, "1.  Where we are today")
# frontend
box(s, 5.17, 1.25, 3.0, 0.55, [("Frontend", 15, WHITE, True)], fill=NAVY)
down(s, 6.67, 1.85)
box(s, 3.35, 2.3, 3.0, 0.6, [("yggdrasil (BFF)", 14, INK, True)], fill=LIGHT, line=GRAY)
box(s, 6.98, 2.3, 3.0, 0.6, [("ems-backend (.NET)", 14, INK, True)], fill=LIGHT, line=GRAY)
down(s, 6.67, 2.95)
box(s, 2.4, 3.45, 8.55, 0.75,
    [("~15 services", 14, INK, True),
     ("energy-model · alarms · reporting · benchmark · export · consumption-api · climate · …", 11, GRAY, False)],
    fill=LIGHT, line=LGRAY)
# fan-out arrows
down(s, 5.0, 4.25); down(s, 8.35, 4.25)
box(s, 3.35, 4.75, 3.0, 0.6, [("analysis_service", 14, WHITE, True)], fill=BLUE)
box(s, 6.98, 4.75, 3.0, 0.6, [("meter_service", 14, WHITE, True)], fill=BLUE)
down(s, 4.85, 5.4); down(s, 8.48, 5.4)
box(s, 3.35, 5.85, 3.0, 0.7, [("counter_measurements", 12, INK, True), ("DB", 10, GRAY, False)],
    fill=WHITE, line=GRAY, shape=MSO_SHAPE.CAN)
box(s, 6.98, 5.85, 3.0, 0.7, [("meter / me2", 12, INK, True), ("DB", 10, GRAY, False)],
    fill=WHITE, line=GRAY, shape=MSO_SHAPE.CAN)
footer(s, "services are tightly coupled around measurements, and the meter model is hard to follow.")

# ============================================================ 2. TARGET
s = slide()
title_bar(s, "2.  Where we want to be")
box(s, 5.17, 1.3, 3.0, 0.55, [("Frontend", 15, WHITE, True)], fill=NAVY)
down(s, 6.67, 1.9)
box(s, 5.17, 2.35, 3.0, 0.55, [("yggdrasil (temporary)", 13, INK, True)], fill=LIGHT, line=GRAY)
down(s, 6.67, 2.95)
box(s, 3.35, 3.4, 3.0, 0.75, [("aggregations", 15, WHITE, True), ("measurement data", 11, LGRAY, False)], fill=GREEN)
box(s, 6.98, 3.4, 3.0, 0.75, [("hierarchy", 15, WHITE, True), ("identity", 11, LGRAY, False)], fill=GREEN)
down(s, 6.67, 4.25)
box(s, 2.4, 4.75, 8.55, 1.15,
    [("Data pipeline", 15, WHITE, True),
     ("Kinesis  →  Flink (enrich + resample)  →  S3/Iceberg (logical_sensor_data)", 12, LGRAY, False),
     ("→  Spark rollup  →  DynamoDB (measurements_aggregate)", 12, LGRAY, False)],
    fill=TEAL)
footer(s, "the frontend reads from one or two services; the pipeline does the computation.")

# ============================================================ 3. EXISTS
s = slide()
title_bar(s, "3.  What already exists")
bullets(s, 0.7, 1.5, 12.0, 4.5, [
    "Spark/Glue rollup: logical_meter_data → DynamoDB view (measurements_aggregate).",
    ("hour + day periods, pre-aggregated at every hierarchy level (company → building → meter).", 1),
    ("group-by is already materialised.", 1),
    "Rust CQRS read API (aggregations) serves it; raw-data reads via Athena for datatilegnelse.",
    "hierarchy service + meter-identity own identity/hierarchy.",
    ("a cross-account CDC bridge already keeps them fed.", 1),
], size=18)
footer(s, "the rollup, the read API and the identity plane already exist; missing = the compute layer and getting all data in.")

# ============================================================ 4. INGESTION
s = slide()
title_bar(s, "4.  Measurements in — ingestion")
chain = ["Sources", "Kinesis\n(EventBridge Pipes)", "Flink\nenrich + resample",
         "logical_sensor_data\n(S3 / Iceberg)", "Spark rollup\n(hour / day)",
         "measurements_aggregate\n(DynamoDB)", "aggregations API\n→ Frontend"]
colors = [LIGHT, LIGHT, TEAL, TEAL, TEAL, GREEN, NAVY]
bw, gap, y = 1.55, 0.3, 2.15
x = 0.35
for i, (label, col) in enumerate(zip(chain, colors)):
    tc = WHITE if col in (TEAL, GREEN, NAVY, BLUE) else INK
    lines = label.split("\n")
    paras = [(lines[0], 12, tc, True)] + [(ln, 10, (LGRAY if tc == WHITE else GRAY), False) for ln in lines[1:]]
    box(s, x, y, bw, 1.5, paras, fill=col, line=(None if col != LIGHT else LGRAY))
    if i < len(chain) - 1:
        right(s, x + bw, y + 0.75, w=gap)
    x += bw + gap
box(s, 0.35, 3.95, 3.4, 1.9,
    [("Source routing", 13, INK, True),
     ("Electrocom → dedicated parser", 11, INK, False),
     ("Catch-all MQTT → AWS IoT Core", 11, INK, False),
     ("Brunata/Datahub/Aalborg/Danfoss/Techem → multi-tenant API", 11, INK, False),
     ("CSV → csv_parser", 11, INK, False),
     ("Manual updates → through the pipe", 11, INK, False),
     ("Kinect → dead (nothing to migrate)", 11, GRAY, False)],
    fill=WHITE, line=LGRAY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
footer(s, "each source is routed into the pipeline; the legacy ingestion services can then be removed.")

# ============================================================ 5. COMPUTE LANES
s = slide()
title_bar(s, "5.  Measurements out — the compute model")
lanes = [
    ("A · Spark rollup", "energy / flow / temperatures — each from its own sensor; cooling & VWAT are derived (flow × Δtemp) via the formula step", BLUE),
    ("B · Flink stream", "active-hours, standby — need sub-hour resolution → daily columns", TEAL),
    ("C · Denormalised columns", "cost = energy×price, CO2 = energy×factor, degree-days from weather — broadcast-joined in Spark on (location, date) / (resource, period)", GREEN),
]
ly = 1.55
for name, desc, col in lanes:
    box(s, 0.6, ly, 5.1, 1.15,
        [(name, 15, WHITE, True), (desc, 11, LGRAY, False)], fill=col, align=PP_ALIGN.LEFT)
    right(s, 5.8, ly + 0.575, w=0.4)
    ly += 1.4
box(s, 6.4, 1.5, 6.35, 4.05, "", fill=LIGHT, line=LGRAY)  # container
textbox(s, 6.62, 1.6, 5.9, 0.32, [("measurements_aggregate  —  one DynamoDB item", 14, INK, True)])
textbox(s, 6.62, 1.96, 5.9, 0.3, [("keyed by:  node · resource · granularity · period", 11, GRAY, False)])
box(s, 6.62, 2.33, 5.9, 0.8,
    [("Columns from the readings  ·  Spark rollup (A)", 11, WHITE, True),
     ("sum · count · min · max · last · temp · flow · cooling", 11, LGRAY, False)],
    fill=BLUE, align=PP_ALIGN.LEFT)
box(s, 6.62, 3.2, 5.9, 0.62,
    [("Columns from fine resolution  ·  Flink (B)", 11, WHITE, True),
     ("active_hours · standby_energy", 11, LGRAY, False)],
    fill=TEAL, align=PP_ALIGN.LEFT)
box(s, 6.62, 3.89, 5.9, 0.8,
    [("Columns from reference data  ·  denormalised (C)", 11, WHITE, True),
     ("cost · co2_scope_* · hdd · cdd · energy_cc", 11, LGRAY, False)],
    fill=GREEN, align=PP_ALIGN.LEFT)
textbox(s, 6.62, 4.78, 5.9, 0.35, [("A read = one DynamoDB lookup — no joins", 12, AMBER, True)])
footer(s, "each lane writes columns onto one aggregate item; a read is a single DynamoDB lookup, no joins.")

# ============================================================ 6. METER MODEL
s = slide()
title_bar(s, "6.  The meter model, simplified")
grid = [("main meter", 0.9, 1.7), ("sub meter", 3.35, 1.7),
        ("calculation meter", 0.9, 2.75), ("\"part of summation\" flag", 3.35, 2.75)]
for label, gx, gy in grid:
    box(s, gx, gy, 2.3, 0.9, [(label, 13, INK, True)], fill=LGRAY, line=GRAY)
textbox(s, 0.9, 3.9, 4.8, 0.5, [("four concepts to avoid double-counting", 12, GRAY, False)])
arrow(s, 6.0, 2.35, 1.3, 0.7, color=TEAL, shape=MSO_SHAPE.RIGHT_ARROW)
box(s, 7.7, 1.9, 4.9, 1.4,
    [("Sensor + formula", 18, WHITE, True),
     ("one concept · Identity / Zero / Expr", 12, LGRAY, False)], fill=GREEN)
bullets(s, 7.7, 3.6, 5.2, 2.4, [
    "sub → main = a single \"part-of-main\" reference; netting is auto-generated.",
    "no second hierarchy — main/sub may cross buildings.",
    "calculation meter → a formula (already stored as one).",
    "users pick from the company tree; formulas are generated, not hand-written.",
], size=13)
box(s, 0.9, 4.5, 5.35, 1.6,
    [("Sensor identity — the new model", 12, INK, True),
     ("Each measure is its own sensor, keyed by daq_id", 11, INK, False),
     ("(protocol:…:meterserial:sensorid) — no counter-number slots.", 11, GRAY, False),
     ("\"Primary\" / temperature come from resource + meter_type,", 11, INK, False),
     ("not TaellerNr (a legacy convention).", 11, GRAY, False)],
    fill=LIGHT, line=LGRAY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
footer(s, "the four meter types become one — a sensor with a formula — with double-counting handled automatically.")

# ============================================================ 7. CONSUMERS
s = slide()
title_bar(s, "7.  Consolidating the consumers — a two-step move")
box(s, 1.2, 1.3, 10.9, 0.95,
    [("Step 1 — Repoint  (migration step)", 14, WHITE, True),
     ("services stop reaching into analysis / meter and read the clean common API — breaks the worst coupling, but they still exist", 11, WHITE, False)],
    fill=AMBER, align=PP_ALIGN.LEFT)
down(s, 6.67, 2.32, h=0.3)
box(s, 1.2, 2.75, 10.9, 0.5, [("Step 2 — Fold in or re-home  (the destination)", 14, WHITE, True)],
    fill=GREEN, align=PP_ALIGN.LEFT)
box(s, 1.2, 3.4, 3.5, 1.75,
    [("Fold into aggregations / hierarchy", 12, WHITE, True), ("", 5, WHITE, False),
     ("measurement / identity logic", 10, LGRAY, False),
     ("consumption-api, computed-benchmark", 10, LGRAY, False)],
    fill=GREEN, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
box(s, 4.9, 3.4, 3.5, 1.75,
    [("Re-home as a bounded-context service", 12, INK, True), ("", 5, INK, False),
     ("genuine separate domains:", 10, GRAY, False),
     ("alarms · reporting / CSRD · ML", 10, INK, False),
     ("own data · loose coupling (events / APIs)", 10, GRAY, False)],
    fill=LIGHT, line=BLUE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
box(s, 8.6, 3.4, 3.5, 1.75,
    [("Retire", 12, INK, True), ("", 5, INK, False),
     ("legacy monolith", 10, GRAY, False),
     ("manual-readings · back-office", 10, GRAY, False),
     ("energy-cost (after we own prices)", 10, GRAY, False)],
    fill=LIGHT, line=GRAY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
box(s, 1.2, 5.4, 10.9, 0.7,
    [("Stopping at Step 1 just relocates the hard coupling — the goal is a small set of context-owning services, not a fan-out of middle-tier ones.", 12, INK, True)],
    fill=WHITE, line=AMBER, line_w=1.5)
footer(s, "repoint first to break the coupling, then fold in or re-home — stopping at repoint would relocate the coupling.")

# ============================================================ 8. VALIDATE
s = slide()
title_bar(s, "8.  How we validate before building")
box(s, 0.7, 1.7, 5.8, 2.6,
    [("Meter-hierarchy test", 16, WHITE, True), ("", 6, WHITE, False),
     ("derive formulas from the old DB", 12, LGRAY, False),
     ("→ recompute totals", 12, LGRAY, False),
     ("→ compare vs current summation, per company", 12, LGRAY, False),
     ("", 6, WHITE, False),
     ("runs offline — no pipeline changes", 12, WHITE, True)],
    fill=BLUE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
box(s, 6.85, 1.7, 5.8, 2.6,
    [("Reference-data test", 16, WHITE, True), ("", 6, WHITE, False),
     ("old values output (cost / CO2 / degree-days)", 12, LGRAY, False),
     ("→ compare vs the new denormalised columns", 12, LGRAY, False)],
    fill=TEAL, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
box(s, 0.7, 4.6, 11.95, 0.9,
    [("Both compare the new model to the current analysis_service output on real data — go/no-go gates before any cutover.", 14, INK, True)],
    fill=LIGHT, line=LGRAY)
footer(s, "we replay the old results and compare; no cutover until the numbers match.")

# ============================================================ 9. PATH
s = slide()
title_bar(s, "9.  The path ahead")
box(s, 0.5, 1.55, 2.7, 1.3,
    [("Phase 0", 15, WHITE, True), ("Verify & size", 12, LGRAY, False), ("(now, read-only)", 10, LGRAY, False)], fill=NAVY)
right(s, 3.25, 2.9, w=0.35)
box(s, 3.7, 1.35, 4.6, 0.9, [("Phase 1 — Extend the data", 13, WHITE, True), ("weather/cost/CO2 + columns", 10, LGRAY, False)], fill=GREEN)
box(s, 3.7, 2.35, 4.6, 0.9, [("Phase 2 — Formula & hierarchy", 13, WHITE, True), ("wire eval; migrate hierarchy", 10, LGRAY, False)], fill=BLUE)
textbox(s, 3.7, 3.28, 4.6, 0.4, [("(the two run in parallel)", 11, GRAY, False)], align=PP_ALIGN.CENTER)
right(s, 8.35, 2.9, w=0.35)
box(s, 8.8, 1.55, 3.9, 1.3,
    [("Phase 3", 15, WHITE, True), ("Consolidate & cut over", 12, LGRAY, False), ("shared switch, once validated", 10, LGRAY, False)], fill=NAVY)
box(s, 0.5, 4.15, 12.2, 0.8,
    [("Track I — Ingestion (parallel foundation):  parsers / EventBridge Pipes / IoT Core so all sources are in the pipeline", 13, INK, True)],
    fill=LIGHT, line=AMBER, line_w=1.5)
box(s, 0.5, 5.2, 12.2, 0.7,
    [("First step:  the meter-hierarchy test — validates the main assumption, runs offline, comes before the largest piece of work.", 13, INK, False)],
    fill=WHITE, line=LGRAY)
footer(s, "validate company by company, build in parallel, then cut over as a shared switch once the numbers match.")

# ============================================================ 10. OPEN
s = slide()
title_bar(s, "10.  Open decisions")
bullets(s, 0.7, 1.6, 12.0, 4.0, [
    "User-facing meter model — auto-netting + one unified concept (recommended) or exposing formulas.",
    "Data-volume sizing — confirm the denormalised DynamoDB item is affordable before schema-freeze.",
    "Own prices — when to retire energy-cost.",
    "Hierarchy migration — mechanical vs dual-run (validated company by company), decided by the test.",
    "Cutover model — a shared switch staged by capability (default), or per-tenant (needs dual ingestion + frontend routing).",
], size=17)
footer(s, "a few decisions remain; the analysis has settled the rest.")

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "counter-measurement-migration.pptx")
prs.save(out)
print("wrote", out, "-", len(prs.slides._sldIdLst), "slides")
